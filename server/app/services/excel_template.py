import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.services.excel_schema import COMMON_COLUMNS, DEAL_SHEETS, GUIDE_TEXT_LINES


def build_template_workbook() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)  # حذف شیت خالی پیش‌فرض

    header_fill = PatternFill("solid", fgColor="2E1065")   # بنفش سورمه‌ای — هماهنگ با تم برنامه
    header_font = Font(bold=True, color="FFFFFF")

    for deal_type, config in DEAL_SHEETS.items():
        ws = wb.create_sheet(title=config["sheet_name"])
        headers = [label for _, label in COMMON_COLUMNS] + [f["header"] for f in config["fields"]]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        # هدر همیشه هنگام اسکرول دیده شود
        ws.freeze_panes = "A2"
        # عرض هر ستون متناسب با طول عنوان
        for col_idx, header in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(32, len(header) + 6))
        ws.sheet_view.rightToLeft = True

    guide_ws = wb.create_sheet(title="راهنما")
    for i, line in enumerate(GUIDE_TEXT_LINES, start=1):
        guide_ws.cell(row=i, column=1, value=line)
    guide_ws.sheet_view.rightToLeft = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()