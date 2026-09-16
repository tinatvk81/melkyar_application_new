import io

from openpyxl import Workbook
from openpyxl.styles import Font

DEAL_TYPE_LABELS_FA = {"sale": "فروش", "presale": "پیش‌خرید", "rent": "اجاره", "mortgage": "رهن کامل"}


def _types_fa(types: list | None) -> str:
    from app.core.property_types import PROPERTY_TYPE_LABELS
    return "، ".join(PROPERTY_TYPE_LABELS.get(t, t) for t in (types or [])) or ""


def build_properties_excel(properties: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "فایل‌ها"
    ws.sheet_view.rightToLeft = True

    headers = [
        "شناسه", "شهر", "منطقه", "آدرس", "متراژ", "اتاق",
        "نوع معامله", "نوع ملک", "وضعیت",
        "قیمت/مبلغ", "قیمت هر متر", "ودیعه", "اجارهٔ ماهانه",
        "آسانسور", "پارکینگ", "امکانات دلخواه",
        "نام مالک", "تلفن مالک",
        "تاریخ پایان قرارداد", "تاریخ ثبت", "آخرین ویرایش",
        "نسخه", "توضیحات",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for p in properties:
        details = p.get("details") or {}
        amenities = p.get("amenities")
        ws.append([
            p.get("id"),
            p.get("city") or "",
            p.get("district") or "",
            p.get("address") or "",
            p.get("area_m2") or "",
            p.get("rooms") or "",
            DEAL_TYPE_LABELS_FA.get(p.get("deal_type"), p.get("deal_type") or ""),
            _types_fa(p.get("property_types")),
            "فعال" if p.get("status") == "active" else ("فروخته‌شده" if p.get("status") == "sold" else "غیرفعال"),
            details.get("price") or details.get("total_price") or details.get("deposit_full") or details.get("monthly_rent") or "",
            details.get("price_per_m2") or "",
            details.get("deposit") or "",
            details.get("monthly_rent") or "",
            "بله" if p.get("has_elevator") else "خیر",
            "بله" if p.get("has_parking") else "خیر",
            "، ".join(amenities) if amenities else "",
            p.get("owner_name") or "",
            p.get("owner_phone") or "",
            p.get("contract_end_date") or "",
            (p.get("created_at") or "")[:19].replace("T", " "),
            (p.get("updated_at") or "")[:19].replace("T", " "),
            p.get("version"),
            p.get("notes") or "",
        ])

    # عرض ستون‌ها
    from openpyxl.utils import get_column_letter
    widths = [8, 14, 14, 34, 8, 7, 12, 20, 11, 16, 14, 14, 14, 9, 9, 24, 16, 15, 14, 18, 18, 7, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()